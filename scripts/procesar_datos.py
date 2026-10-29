# /// script
# requires-python = ">=3.11"
# dependencies = ["pandas>=2.2", "openpyxl>=3.1", "python-calamine>=0.2", "duckdb>=1.1"]
# ///
"""Convierte los archivos de data_raw/ en los archivos livianos que usa el sitio.

Uso (desde la raíz del repositorio; descargar_datos.py actualiza antes los mercados y las
previsiones del FMI y del Banco Mundial):

    uv run scripts/descargar_datos.py
    uv run scripts/procesar_datos.py

Salidas en src/data/:
    series.csv / catalogo.json     series macro en formato largo (id, fecha, valor)
    mercados.csv / mercados.json   series diarias de mercados internacionales en formato ancho
    previsiones.csv                previsiones de crecimiento del PIB por institución
    exportaciones.csv / .json      exportaciones mensuales por producto (clasificación BCE)
    importaciones.csv / .json      importaciones mensuales por CUODE
    ventas_sector.csv              ventas locales SRI mensuales por sección CIIU (y división en comercio)
    ipc_incidencias.csv            incidencia mensual por división del IPC (último mes)
    spnf.csv                       operaciones mensuales del SPNF por partida (MEF)
    vencimientos.csv               perfil de vencimientos de la deuda pública (MEF)
"""

from __future__ import annotations

import json
import re
from pathlib import Path

import duckdb
import pandas as pd

from catalogo_mercados import CURVAS, FUENTES, MERCADOS

RAIZ = Path(__file__).resolve().parent.parent
RAW = RAIZ / "data_raw"
OUT = RAIZ / "src" / "data"

MESES = {m: i + 1 for i, m in enumerate(
    ["ene", "feb", "mar", "abr", "may", "jun", "jul", "ago", "sep", "oct", "nov", "dic"])}
MESES_LARGOS = {m: i + 1 for i, m in enumerate(
    ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto",
     "septiembre", "octubre", "noviembre", "diciembre"])}
ROMANOS = {"I": 1, "II": 4, "III": 7, "IV": 10}

filas: list[tuple[str, str, float]] = []
catalogo: list[dict] = []


def limpiar(texto) -> str:
    return " ".join(str(texto).split())


def registrar(id_, fechas, valores, *, sector, nombre, corto, unidad, frecuencia,
              decimales, fuente, escala=1.0):
    """Agrega una serie (descartando vacíos) y su ficha de catálogo."""
    n = 0
    for f, v in zip(fechas, valores):
        if v is None or pd.isna(v):
            continue
        filas.append((id_, pd.Timestamp(f).strftime("%Y-%m-%d"), round(float(v) * escala, 6)))
        n += 1
    if n == 0:
        raise ValueError(f"Serie vacía: {id_}")
    catalogo.append(dict(id=id_, sector=sector, nombre=nombre, corto=corto, unidad=unidad,
                         frecuencia=frecuencia, decimales=decimales, fuente=fuente))


# --------------------------------------------------------------- utilidades de lectura
def hoja_cn(ruta: Path, hoja: str) -> pd.DataFrame:
    """Hoja de cuentas nacionales trimestrales (filas '2026.II') → DataFrame por fecha."""
    df = pd.read_excel(ruta, sheet_name=hoja, header=None)
    fila_cab = next(i for i in range(20) if limpiar(df.iat[i, 0]) in ("Industrias", "Variables"))
    cab = [limpiar(c) for c in df.iloc[fila_cab]]
    datos = df.iloc[fila_cab + 1:].copy()
    datos.columns = cab
    etiqueta = datos.iloc[:, 0].astype(str).str.strip()
    datos = datos[etiqueta.str.match(r"^\d{4}\.(I|II|III|IV)$")]
    datos.index = [pd.Timestamp(int(t[:4]), ROMANOS[t[5:]], 1) for t in datos.iloc[:, 0].str.strip()]
    return datos.iloc[:, 1:].apply(pd.to_numeric, errors="coerce")


def mas_reciente(carpeta: Path, patron: str) -> Path:
    """Último archivo de una serie de publicaciones (el modificado más recientemente)."""
    rutas = list(carpeta.glob(patron))
    if not rutas:
        raise FileNotFoundError(f"No hay archivos {patron} en {carpeta}")
    return max(rutas, key=lambda r: r.stat().st_mtime)


def hoja_imaec(hoja: str) -> pd.DataFrame:
    df = pd.read_excel(mas_reciente(RAW / "IMAEc", "IMAEc_*.xlsx"), sheet_name=hoja, header=None)
    fila_cab = next(i for i in range(20) if limpiar(df.iat[i, 0]) == "Fecha")
    datos = df.iloc[fila_cab + 1:].copy()
    datos.columns = [limpiar(c) for c in df.iloc[fila_cab]]
    datos = datos[datos["Fecha"].astype(str).str.match(r"^\d{4}\.[a-z]{3}$")]
    datos.index = [pd.Timestamp(int(f[:4]), MESES[f[5:]], 1) for f in datos["Fecha"]]
    return datos.drop(columns="Fecha").apply(pd.to_numeric, errors="coerce")


def cuadro_bce(ruta: Path) -> dict[str, pd.DataFrame]:
    """Cuadro de BCEData (series en filas, meses en columnas) → {bloque: DataFrame por fecha}.

    Un cuadro trae uno o más bloques, cada uno con su fila 'Serie' de fechas ('Ene 2006').
    Las filas sin valores son subtítulos ('Variación Mensual…', 'Segmento Consumo'): dan nombre
    al bloque que sigue. Los bloques que no son mensuales (p. ej. semanales) se omiten.
    """
    df = pd.read_excel(ruta, header=None, dtype=str)
    bloques, actual, titulo, fechas = {}, None, "", None
    for i in range(len(df)):
        etiqueta = limpiar(df.iat[i, 0]) if pd.notna(df.iat[i, 0]) else ""
        valores = df.iloc[i, 1:]
        if etiqueta == "Serie":
            cols = {j: c for j, c in valores.items() if isinstance(c, str) and re.fullmatch(r"[A-Z][a-z]{2} \d{4}", c)}
            fechas = {j: pd.Timestamp(int(c[4:]), MESES[c[:3].lower()], 1) for j, c in cols.items()} or None
            titulo, actual = "", None
        elif not etiqueta or etiqueta.startswith(("Unidad", "Frecuencia")):
            fechas = None if etiqueta.startswith("Frecuencia") else fechas
        elif fechas and valores.isna().all():
            titulo, actual = etiqueta, None
        elif fechas:
            if actual is None:
                actual = bloques.setdefault(titulo, {})
            actual[etiqueta] = pd.Series(
                [float(str(valores[j]).replace(",", ".")) if pd.notna(valores[j]) else None for j in fechas],
                index=list(fechas.values()), dtype=float)
    return {t: pd.DataFrame(s) for t, s in bloques.items()}


def slug(texto: str) -> str:
    import unicodedata
    plano = unicodedata.normalize("NFKD", texto).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]+", "_", plano).strip("_")


# ============================================================== SECTOR REAL
FUENTE_CN = "BCE, Cuentas Nacionales Trimestrales"
cn = RAW / "cuentas_nacionales_trim"

dem_vy = hoja_cn(cn / "tou_136_202602.xlsx", "Dem_Ivol_ajus_vY")
registrar("pib_real_aa", dem_vy.index, dem_vy["PIB"], sector="real",
          nombre="PIB real, variación anual (serie desestacionalizada)", corto="Crecimiento del PIB",
          unidad="%", frecuencia="T", decimales=1, fuente=FUENTE_CN)

# Aportes al crecimiento anual del PIB por componente del gasto.
contr = hoja_cn(cn / "tou_136_202602.xlsx", "Dem_contr_ajus_vY")
col = {c.lower(): c for c in contr.columns}
buscar = lambda texto: next(c for k, c in col.items() if texto in k)
c_hog, c_gob = buscar("consumo final hogar"), buscar("consumo final gobierno")
c_fbkf, c_exist = buscar("fbkf"), buscar("existencias")
c_x, c_m = buscar("exportaciones"), buscar("importaciones")
netas = contr[c_x] - contr[c_m]
residuo = (contr[c_hog] + contr[c_gob] + contr[c_fbkf] + contr[c_exist] + netas - contr["PIB"]).abs().max()
if residuo > 0.2:  # si el signo de las importaciones viniera invertido
    netas = contr[c_x] + contr[c_m]
for id_, valores, corto in [
    ("aporte_consumo", contr[c_hog] + contr[c_gob], "Consumo"),
    ("aporte_fbkf", contr[c_fbkf], "Inversión (FBKF)"),
    ("aporte_xnetas", netas, "Exportaciones netas"),
    ("aporte_existencias", contr[c_exist], "Variación de existencias"),
]:
    registrar(id_, contr.index, valores, sector="real", nombre=f"Aporte al crecimiento anual del PIB: {corto}",
              corto=corto, unidad="pp", frecuencia="T", decimales=2, fuente=FUENTE_CN)

dem_corr = hoja_cn(cn / "tou_136_202602.xlsx", "Dem_Corr_bru")
pib_nominal_t = dem_corr["PIB"] / 1000  # miles → millones de USD
pib_anual = pib_nominal_t.groupby(pib_nominal_t.index.year).agg(["sum", "count"])
pib_anual = pib_anual[pib_anual["count"] == 4]["sum"]
registrar("pib_nominal", [pd.Timestamp(a, 1, 1) for a in pib_anual.index], pib_anual, sector="real",
          nombre="PIB nominal", corto="PIB nominal", unidad="USD mm", frecuencia="A", decimales=0,
          fuente=FUENTE_CN)

pet = hoja_cn(cn / "vab_p_np_136_202602.xlsx", "Of_Ivol_ajus_pet_vY")
registrar("vab_petrolero_aa", pet.index, pet["Valor Agregado Petrolero"], sector="real",
          nombre="VAB petrolero real, variación anual", corto="VAB petrolero", unidad="%",
          frecuencia="T", decimales=1, fuente=FUENTE_CN)
registrar("vab_no_petrolero_aa", pet.index, pet["Valor Agregado no Petrolero"], sector="real",
          nombre="VAB no petrolero real, variación anual", corto="VAB no petrolero", unidad="%",
          frecuencia="T", decimales=1, fuente=FUENTE_CN)

# IMAEc
FUENTE_IMAE = "BCE, IMAEc"
imae = hoja_imaec("Of_Ivol_ajus_clas")
imae_vy = hoja_imaec("Of_Ivol_ajus_clas_vY")
registrar("imaec", imae.index, imae["IMAEc"], sector="real",
          nombre="Índice de Actividad Económica coyuntural (IMAEc), desestacionalizado",
          corto="IMAEc", unidad="índice", frecuencia="M", decimales=1, fuente=FUENTE_IMAE)
registrar("imaec_aa", imae_vy.index, imae_vy["IMAEc"], sector="real",
          nombre="IMAEc, variación anual", corto="IMAEc, variación anual", unidad="%",
          frecuencia="M", decimales=1, fuente=FUENTE_IMAE)
for col_, id_, corto in [
    ("Agricultura", "agro", "Agropecuario"), ("Petróleo", "petroleo", "Petróleo y minas"),
    ("Manufactura", "manufactura", "Manufactura"), ("Construcción", "construccion", "Construcción"),
    ("Comercio", "comercio", "Comercio"), ("Servicios", "servicios", "Servicios"),
]:
    c = next(c for c in imae_vy.columns if c.startswith(col_))
    registrar(f"imaec_aa_{id_}", imae_vy.index, imae_vy[c], sector="real",
              nombre=f"IMAEc {corto}, variación anual", corto=corto, unidad="%", frecuencia="M",
              decimales=1, fuente=FUENTE_IMAE)

# Inflación: cuadro 4.2.1 del BCE (índice y variaciones del IPC general y por división).
FUENTE_IPC = "BCE, con datos del INEC (Índice de Precios al Consumidor)"
ipc = cuadro_bce(RAW / "IPC BCE.xlsx")
GENERAL = "Índice General - IPC"
for bloque, id_, nombre, corto in [
    ("Variación Anual del Índice de Precios al Consumidor", "inflacion_aa", "Inflación anual (IPC)", "Inflación anual"),
    ("Variación Mensual del Índice de Precios al Consumidor", "inflacion_mm", "Inflación mensual (IPC)", "Inflación mensual"),
    ("Variación Acumulada del Índice de Precios al Consumidor", "inflacion_acum", "Inflación acumulada en el año (IPC)",
     "Inflación acumulada"),
]:
    s = ipc[bloque][GENERAL].dropna()
    s = s[s.index >= "2010-01-01"]
    registrar(id_, s.index, s, sector="real", nombre=nombre, corto=corto, unidad="%",
              frecuencia="M", decimales=2, fuente=FUENTE_IPC)

# Incidencia por división en el último mes (índice de Laspeyres):
#   incidencia_i = ponderación_i × (índice_i / índice general)_{t−1} × variación mensual_i,t
# Las ponderaciones de la canasta vienen del cuadro de incidencias del INEC; las divisiones
# están en el mismo orden (CCIF 01 a 13) en ambos archivos.
pond = pd.read_excel(RAW / "IPC nacional.xlsx", sheet_name="1.Inc_mens", header=7)
pond = pond[pond["Nivel"] == "División"]
indices = ipc[""]
var_mm = ipc["Variación Mensual del Índice de Precios al Consumidor"]
divisiones = [c for c in indices.columns if c != GENERAL]
assert len(divisiones) == len(pond) == 13, "Las divisiones del IPC del BCE y del INEC no coinciden"
fecha_incidencias = var_mm[GENERAL].dropna().index.max()
previo_ipc = indices.loc[fecha_incidencias - pd.DateOffset(months=1)]
incidencia = [p * previo_ipc[d] / previo_ipc[GENERAL] * var_mm.at[fecha_incidencias, d]
              for p, d in zip(pond["Ponderación"], divisiones)]
residuo = abs(sum(incidencia) - var_mm.at[fecha_incidencias, GENERAL])
assert residuo < 0.01, f"Las incidencias del IPC no suman la inflación mensual (diferencia {residuo:.4f})"
pd.DataFrame({
    "ccif": pond["CCIF"].astype(str).str.zfill(2).values,
    "division": pond["Descripción CCIF"].map(limpiar).values,
    "ponderacion": (pond["Ponderación"] * 100).round(3).values,
    "incidencia": [round(x, 4) for x in incidencia],
}).to_csv(OUT / "ipc_incidencias.csv", index=False)

# Mercado laboral (microdatos ENEMDU, ponderados con el factor de expansión)
con = duckdb.connect()
enemdu = con.execute(f"""
    select cast(periodo as int) periodo,
      sum(fexp) filter (where condact between 1 and 8) pea,
      sum(fexp) filter (where condact in (7, 8)) desempleo,
      sum(fexp) filter (where condact = 1) adecuado,
      sum(fexp) filter (where condact in (2, 3)) subempleo,
      sum(fexp) filter (where condact in (4, 5, 6)) otro_no_pleno
    from '{RAW / "enemdu_mensual.parquet"}' group by 1 order by 1
""").df()
enemdu["fecha"] = [pd.Timestamp(p // 100, p % 100, 1) for p in enemdu["periodo"]]
for col_, id_, nombre, corto in [
    ("desempleo", "desempleo", "Tasa de desempleo nacional", "Desempleo"),
    ("adecuado", "empleo_adecuado", "Tasa de empleo adecuado", "Empleo adecuado"),
    ("subempleo", "subempleo", "Tasa de subempleo", "Subempleo"),
    ("otro_no_pleno", "otro_no_pleno", "Otro empleo no pleno y no remunerado", "Otro no pleno"),
]:
    registrar(id_, enemdu["fecha"], enemdu[col_] / enemdu["pea"] * 100, sector="real", nombre=nombre,
              corto=corto, unidad="%", frecuencia="M", decimales=1,
              fuente="INEC, ENEMDU (cálculo propio con factor de expansión)")

# Coyuntura
iee = pd.read_excel(RAW / "coyuntura" / "IEE.xlsx", header=7)
iee = iee[pd.to_datetime(iee["Fecha"], errors="coerce").notna()]
fechas = pd.to_datetime(iee["Fecha"])
for col_, id_, corto in [("IEE Global (2)", "iee", "IEE global"), ("Comercio", "iee_comercio", "Comercio"),
                         ("Construcción", "iee_construccion", "Construcción"),
                         ("Manufactura", "iee_manufactura", "Manufactura"), ("Servicios", "iee_servicios", "Servicios")]:
    registrar(id_, fechas, pd.to_numeric(iee[col_]), sector="real",
              nombre=f"Índice de Expectativas Empresariales: {corto}", corto=corto, unidad="índice",
              frecuencia="M", decimales=1, fuente="Por confirmar")

icc = pd.read_excel(RAW / "coyuntura" / "ICC.xlsx", header=7)
icc = icc[icc["Periodo"].astype(str).str.match(r"^\d{4}-\d{2}$")]
fechas = pd.to_datetime(icc["Periodo"] + "-01")
for col_, id_, corto in [("ICC Global (2)", "icc", "ICC global"),
                         ("Índice de Situación Presente", "icc_presente", "Situación presente"),
                         ("Índice de Situación Futura", "icc_futura", "Situación futura")]:
    c = next(x for x in icc.columns if str(x).startswith(col_))
    registrar(id_, fechas, pd.to_numeric(icc[c]), sector="real",
              nombre=f"Índice de Confianza del Consumidor: {corto}", corto=corto, unidad="índice",
              frecuencia="M", decimales=1, fuente="Por confirmar")

cem = pd.read_excel(RAW / "coyuntura" / "EstadisticasCemento.xlsx", sheet_name="Producción_despachos_inven",
                    header=None)
anio, fechas, despachos, produccion = None, [], [], []
for _, r in cem.iterrows():
    etiqueta = str(r[1]).strip()
    m = re.match(r"^(\d{4})\s+(\w+)$", etiqueta)
    if m:
        anio, mes = int(m[1]), m[2].lower()
    elif anio and etiqueta.lower() in MESES_LARGOS:
        mes = etiqueta.lower()
    else:
        continue
    fechas.append(pd.Timestamp(anio, MESES_LARGOS[mes], 1))
    produccion.append(r[2])
    despachos.append(r[3])
registrar("cemento_despachos", fechas, despachos, sector="real", nombre="Despachos de cemento",
          corto="Despachos de cemento", unidad="miles t", frecuencia="M", decimales=1,
          fuente="Empresas cementeras", escala=1 / 1000)
registrar("cemento_produccion", fechas, produccion, sector="real", nombre="Producción de cemento",
          corto="Producción de cemento", unidad="miles t", frecuencia="M", decimales=1,
          fuente="Empresas cementeras", escala=1 / 1000)

# Ventas SRI
SECCIONES = {
    "A": "Agricultura y pesca", "B": "Minas y canteras", "C": "Manufactura", "D": "Electricidad y gas",
    "E": "Agua y saneamiento", "F": "Construcción", "G": "Comercio", "H": "Transporte y almacenamiento",
    "I": "Alojamiento y comidas", "J": "Información y comunicación", "K": "Actividades financieras",
    "L": "Actividades inmobiliarias", "M": "Actividades profesionales", "N": "Servicios administrativos",
    "O": "Administración pública", "P": "Enseñanza", "Q": "Salud", "R": "Artes y recreación",
    "S": "Otros servicios", "T": "Hogares", "U": "Organizaciones extraterritoriales",
}
# Ventas locales: total de ventas y exportaciones (casillero 419) menos exportaciones de bienes (417)
# y de servicios (418). Por sección CIIU y, en comercio, por división (G45 vehículos, G46 al por mayor,
# G47 al por menor).
DIVISIONES_COMERCIO = {"G45": "Vehículos y motocicletas", "G46": "Comercio al por mayor",
                       "G47": "Comercio al por menor"}
ventas = con.execute(f"""
    select cast("ANIO FISCAL" as int) anio, cast("MES FISCAL" as int) mes,
           upper(left("ACTIVIDAD ECONÓMICA", 3)) division,
           sum("TOTAL VENTAS Y EXPORTACIONES (419)" - coalesce("EXPORTACIONES DE BIENES (417)", 0)
               - coalesce("EXPORTACIONES DE SERVICIOS (418)", 0)) / 1e6 ventas
    from '{RAW / "ventas-saiku-ciiu6.parquet"}' group by all
""").df()
ventas["fecha"] = [pd.Timestamp(a, m, 1) for a, m in zip(ventas["anio"], ventas["mes"])]
ventas["seccion"] = ventas["division"].str[0]
por_seccion = ventas.groupby(["fecha", "seccion"], as_index=False)["ventas"].sum()
por_seccion["codigo"] = por_seccion["seccion"]
por_seccion["nombre"] = por_seccion["seccion"].map(SECCIONES).fillna("Sin clasificar")
comercio = ventas[ventas["division"].isin(DIVISIONES_COMERCIO)].groupby(["fecha", "division"], as_index=False)["ventas"].sum()
comercio["codigo"] = comercio["division"]
comercio["nombre"] = comercio["division"].map(DIVISIONES_COMERCIO)
pd.concat([por_seccion, comercio])[["fecha", "codigo", "nombre", "ventas"]].sort_values(["fecha", "codigo"]).assign(
    fecha=lambda d: d["fecha"].dt.strftime("%Y-%m-%d"), ventas=lambda d: d["ventas"].round(3)
).to_csv(OUT / "ventas_sector.csv", index=False)
total_ventas = por_seccion.groupby("fecha")["ventas"].sum()
registrar("ventas_locales", total_ventas.index, total_ventas, sector="real",
          nombre="Ventas locales declaradas (formulario 104: casillero 419 menos exportaciones, 417 y 418)",
          corto="Ventas locales", unidad="USD mm", frecuencia="M", decimales=0, fuente="SRI, Saiku")

# ============================================================== SECTOR FINANCIERO
# Sistema de bancos privados: indicadores internos con datos de la Superintendencia de Bancos.
# Un panel por fecha y banco (miles de USD y razones); el total del sistema es la fila del
# segmento "BANCOS PRIVADOS" sin banco. Los ceros marcan meses sin dato (p. ej. la solvencia
# antes de 2020 o los segmentos de crédito antes de su creación).
FUENTE_SB = "Superintendencia de Bancos (compilación interna)"
ib = pd.read_excel(mas_reciente(RAW / "sb", "indicadores_banco*.xlsx"), sheet_name="Sheet1", engine="calamine")
ib = ib[(ib["Segmento"] == "BANCOS PRIVADOS") & ib["Banco"].isna()]
ib = ib.set_index(pd.to_datetime(ib["Fecha"]).dt.to_period("M").dt.to_timestamp()).sort_index()
ib = ib.apply(pd.to_numeric, errors="coerce").replace(0, float("nan"))
ib["ltd"] = ib["Cartera Bruta Total"] / ib["Obligaciones Pub"]

SEGMENTOS_SB = [("Productivo", "productivo", "Productivo"), ("Consumo", "consumo", "Consumo"),
                ("Inmobiliario", "inmobiliario", "Inmobiliario"), ("Microcredito", "microcredito", "Microcrédito")]
for col_, id_, nombre, corto, unidad, esc, dec in [
    ("Obligaciones Pub", "depositos_bp", "Depósitos del público, bancos privados", "Depósitos", "USD mm", 1e-3, 0),
    ("Depositos Monetarios", "dep_monetarios_bp", "Depósitos monetarios, bancos privados", "Monetarios", "USD mm", 1e-3, 0),
    ("Depositos Ahorro", "dep_ahorro_bp", "Depósitos de ahorro, bancos privados", "Ahorro", "USD mm", 1e-3, 0),
    ("Depositos Plazo", "dep_plazo_bp", "Depósitos a plazo, bancos privados", "Plazo", "USD mm", 1e-3, 0),
    ("Cartera Bruta Total", "cartera_bp", "Cartera bruta, bancos privados", "Cartera", "USD mm", 1e-3, 0),
    *[(f"Cartera Bruta {c}", f"cartera_{i}_bp", f"Cartera bruta {n.lower()}, bancos privados", n, "USD mm", 1e-3, 0)
      for c, i, n in SEGMENTOS_SB],
    ("Activos", "activos_bp", "Activos totales, bancos privados", "Activos", "USD mm", 1e-3, 0),
    ("ltd", "ltd_bp", "Relación cartera / depósitos (LTD), bancos privados", "Cartera / depósitos", "%", 100, 1),
    ("Proporcion Vista", "dep_vista_pct", "Depósitos a la vista, % del total, bancos privados",
     "Depósitos a la vista", "%", 100, 1),
    ("Morosidad Total", "morosidad_bp", "Morosidad de la cartera (improductiva / bruta), bancos privados",
     "Morosidad", "%", 100, 2),
    *[(f"Morosidad {c}", f"morosidad_{i}_bp", f"Morosidad de la cartera {n.lower()}, bancos privados", n, "%", 100, 2)
      for c, i, n in SEGMENTOS_SB],
    ("Cobertura Total", "cobertura_bp", "Cobertura de la cartera improductiva (provisiones / improductiva), bancos privados",
     "Cobertura", "%", 100, 0),
    ("Pt Activos Contingente Riesgo", "solvencia_bp",
     "Solvencia (patrimonio técnico / activos y contingentes ponderados por riesgo), bancos privados",
     "Solvencia", "%", 100, 1),
    # Primera y segunda línea: aproximaciones con información pública a los indicadores regulatorios.
    ("Liquidez Primera Linea", "liquidez_primera_bp",
     "Liquidez de primera línea (aproximación con información pública), bancos privados", "Primera línea", "%", 100, 1),
    ("Liquidez Segunda Linea", "liquidez_segunda_bp",
     "Liquidez de segunda línea (aproximación con información pública), bancos privados", "Segunda línea", "%", 100, 1),
    ("Roe", "roe_bp", "Rentabilidad sobre patrimonio (ROE), bancos privados", "ROE", "%", 100, 1),
    ("Camel Roa", "roa_bp", "Rentabilidad sobre activos (ROA), bancos privados", "ROA", "%", 100, 2),
]:
    registrar(id_, ib.index, ib[col_], sector="financiero", nombre=nombre, corto=corto, unidad=unidad,
              frecuencia="M", decimales=dec, fuente=FUENTE_SB, escala=esc)

# Liquidez oficial de la SB (la que reporta Asobanca): fondos disponibles / total de depósitos a
# corto plazo (a la vista y a plazo hasta 90 días), del archivo mensual de series del subsistema
# de bancos privados (hoja INDICAD, una columna por mes).
# El nombre termina en AAAA_MM: se usa el último mes.
series_sb = pd.read_excel(max((RAW / "sb").glob("SERIES SUBSISTEMA BANCA PRIVADA*.xlsx")), sheet_name="INDICAD",
                          header=None, engine="calamine")
i = next(i for i in range(len(series_sb))
         if limpiar(series_sb.iat[i, 0]).upper() == "FONDOS DISPONIBLES / TOTAL DEPOSITOS A CORTO PLAZO")
k = next(k for k in range(i, -1, -1) if sum(hasattr(c, "year") for c in series_sb.iloc[k]) > 3)
cols = [j for j, c in enumerate(series_sb.iloc[k]) if hasattr(c, "year")]
registrar("liquidez_bp", [pd.Timestamp(series_sb.iat[k, j]).to_period("M").to_timestamp() for j in cols],
          [series_sb.iat[i, j] for j in cols], sector="financiero",
          nombre="Fondos disponibles / depósitos a corto plazo (no incluye otros activos líquidos), bancos privados",
          corto="Fondos disponibles / depósitos CP",
          unidad="%", frecuencia="M", decimales=2, escala=100,
          fuente="Superintendencia de Bancos, Series del subsistema de bancos privados")

# Tasas de interés (BCE): referenciales mensuales y activas efectivas por segmento de crédito.
ref = cuadro_bce(RAW / "tasas" / "tasas_referenciales.xlsx")[""]
for col_, id_, nombre, corto in [
    ("Pasiva", "tasa_pasiva", "Tasa pasiva referencial", "Tasa pasiva referencial"),
    ("Activa", "tasa_activa", "Tasa activa referencial", "Tasa activa referencial"),
]:
    registrar(id_, ref.index, ref[col_], sector="financiero", nombre=nombre, corto=corto, unidad="%",
              frecuencia="M", decimales=2, fuente="BCE, Tasas de interés referenciales")

# Cada segmento trae la tasa efectiva referencial y la máxima ("Máxima Consumo"):
# id tasa_<subsegmento> y tasa_max_<subsegmento>.
for df in cuadro_bce(RAW / "tasas" / "tasas_activas_efectivas_maximas.xlsx").values():
    for col_ in df.columns:
        maxima = col_.startswith("Máxima ")
        sub = col_.removeprefix("Máxima ")
        sub = sub[:1] + sub[1:].lower().replace("pymes", "PYMES")  # "Productivo Corporativo" → "Productivo corporativo"
        registrar(f"tasa_{'max_' if maxima else ''}{slug(sub)}", df.index, df[col_], sector="financiero",
                  nombre=f"Tasa activa {'máxima' if maxima else 'efectiva referencial'}: {sub}",
                  corto=sub, unidad="%", frecuencia="M", decimales=2,
                  fuente="BCE, Tasas de interés activas efectivas referenciales y máximas")

# Riesgo país (descargar_datos.py): el BCE repite a veces el viernes en el fin de semana; se dejan días hábiles.
embi = pd.DataFrame(json.loads((RAW / "riesgo_pais.json").read_text())["view_ind_formulario"])
embi = embi[embi["Indicador"] == "Riesgo País"]
embi["Fecha"] = pd.to_datetime(embi["Fecha"])
embi = embi[embi["Fecha"].dt.dayofweek < 5]
registrar("embi", embi["Fecha"], embi["Valor"], sector="financiero", nombre="Riesgo país (EMBI Ecuador)",
          corto="Riesgo país", unidad="pb", frecuencia="D", decimales=0, fuente="BCE, con datos de JP Morgan")

# Bonos globales: precios del BCE (descargar_datos.py); como en el riesgo país, se dejan días hábiles.
bonos = pd.DataFrame(json.loads((RAW / "bonos_soberanos.json").read_text())["view_ind_bonos"])
bonos["Fecha"] = pd.to_datetime(bonos["Fecha"])
bonos = bonos[bonos["Fecha"].dt.dayofweek < 5]
for vencimiento, df in bonos.groupby("Indicador"):
    df = df.sort_values("Fecha")
    registrar(f"bono_{vencimiento}", df["Fecha"], df["Valor"], sector="financiero",
              nombre=f"Bono global Ecuador {vencimiento}: precio", corto=f"Ecuador {vencimiento}",
              unidad="USD por 100", frecuencia="D", decimales=2, fuente="BCE, Bonos soberanos")

# Reservas internacionales semanales. Códigos del BCE:
#   c1 posición neta en divisas (c11 caja en divisas; c12 depósitos netos en bancos e
#      instituciones financieras del exterior; c13 inversiones, depósitos a plazo y títulos)
#   c2 oro; c3 DEG; c4 posición de reserva en el FMI; c5 posición en ALADI; c6 posición en SUCRE
# El resto (c3 a c6) se calcula como RI − c1 − c2: desde febrero de 2026 la fuente no trae c3 (DEG),
# aunque el total sí lo incluye.
ri = con.execute(f"""
    select FECHA fecha, RI ri, c1 divisas, c2 oro, RI - coalesce(c1, 0) - coalesce(c2, 0) otros
    from '{RAW / "reserva_internacional_semanal.parquet"}' order by 1
""").df()
FUENTE_RI = "BCE, Reservas internacionales (semanal)"
for id_, valores, nombre, corto in [
    ("reservas", ri["ri"], "Reservas internacionales", "Reservas internacionales"),
    ("reservas_divisas", ri["divisas"], "Reservas internacionales: posición neta en divisas",
     "Posición neta en divisas"),
    ("reservas_oro", ri["oro"], "Reservas internacionales: oro", "Oro"),
    ("reservas_otros", ri["otros"], "Reservas internacionales: DEG, FMI, ALADI y SUCRE",
     "DEG, FMI, ALADI y SUCRE"),
]:
    registrar(id_, ri["fecha"], valores, sector="financiero", nombre=nombre, corto=corto, unidad="USD mm",
              frecuencia="S", decimales=0, fuente=FUENTE_RI)

# ============================================================== SECTOR FISCAL
import warnings  # noqa: E402

warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")

# Operaciones del SPNF (MEF): columnas anuales, trimestrales y mensuales; se usan las mensuales.
FUENTE_SPNF = "MEF, Operaciones del Sector Público No Financiero"
op = pd.read_excel(mas_reciente(RAW / "mef_spnf", "*Operaciones-de-Ingresos-y-Gastos-SPNF*.xlsx"),
                   sheet_name="SPNF", header=None)
fila_fechas = op.iloc[4]
cols_mes = [j for j, v in enumerate(fila_fechas) if isinstance(v, pd.Timestamp) or hasattr(v, "month")]
fechas_spnf = [pd.Timestamp(fila_fechas[j]) for j in cols_mes]
spnf = []
for i in range(5, len(op)):
    codigo, concepto = op.iat[i, 0], op.iat[i, 1]
    if pd.isna(codigo) or pd.isna(concepto):
        continue
    codigo = str(codigo).split()[0].split(".")[0]
    for j, f in zip(cols_mes, fechas_spnf):
        v = pd.to_numeric(op.iat[i, j], errors="coerce")
        if pd.notna(v):
            spnf.append((f.strftime("%Y-%m-%d"), codigo, limpiar(concepto), round(float(v), 3)))
spnf = pd.DataFrame(spnf, columns=["fecha", "codigo", "concepto", "valor"]).drop_duplicates(["fecha", "codigo"])
# Quitar notas al pie del nombre ("Exportación 1/" → "Exportación").
spnf["concepto"] = spnf["concepto"].str.replace(r"\s*\d+/\s*$|\s*\(\*+\)\s*$|\s*/\d+$", "", regex=True)
spnf.to_csv(OUT / "spnf.csv", index=False)

por_codigo = spnf.pivot(index="fecha", columns="codigo", values="valor")
por_codigo.index = pd.to_datetime(por_codigo.index)
for codigo, id_, nombre, corto in [
    ("1", "spnf_ingresos", "Ingresos totales del SPNF", "Ingresos"),
    ("2", "spnf_gastos", "Gastos totales del SPNF", "Gastos"),
    ("3", "spnf_resultado_global", "Resultado global del SPNF", "Resultado global"),
    ("4", "spnf_resultado_primario", "Resultado primario del SPNF", "Resultado primario"),
    ("41", "spnf_primario_no_petrolero", "Resultado primario no petrolero del SPNF", "Primario no petrolero"),
    ("11", "spnf_ing_petroleros", "Ingresos petroleros del SPNF", "Petroleros"),
    ("121", "spnf_tributarios", "Ingresos tributarios del SPNF", "Tributarios"),
    ("213", "spnf_intereses", "Intereses de la deuda pública", "Intereses"),
    ("221", "spnf_inversion", "Inversión en activos no financieros", "Inversión"),
    ("57", "spnf_subsidios_combustibles", "Subsidios a los combustibles", "Subsidios a combustibles"),
]:
    registrar(id_, por_codigo.index, por_codigo[codigo], sector="fiscal", nombre=nombre, corto=corto,
              unidad="USD mm", frecuencia="M", decimales=0, fuente=FUENTE_SPNF)

# Resultados anuales en % del PIB (años completos), con el PIB nominal de cuentas nacionales.
anual = por_codigo.groupby(por_codigo.index.year).agg(["sum", "count"])
for codigo, id_, nombre, corto in [
    ("3", "resultado_global_pib", "Resultado global del SPNF", "Resultado global"),
    ("4", "resultado_primario_pib", "Resultado primario del SPNF", "Resultado primario"),
    ("41", "primario_no_petrolero_pib", "Resultado primario no petrolero del SPNF", "Primario no petrolero"),
    ("1", "ingresos_spnf_pib", "Ingresos totales del SPNF", "Ingresos"),
    ("2", "gastos_spnf_pib", "Gastos totales del SPNF", "Gastos"),
]:
    completos = [a for a in anual.index if anual.loc[a, (codigo, "count")] == 12 and a in pib_anual.index]
    registrar(id_, [pd.Timestamp(a, 1, 1) for a in completos],
              [anual.loc[a, (codigo, "sum")] / pib_anual[a] * 100 for a in completos], sector="fiscal",
              nombre=nombre, corto=corto, unidad="% PIB", frecuencia="A", decimales=1,
              fuente=f"{FUENTE_SPNF}; BCE, Cuentas Nacionales (cálculo propio)")

# Deuda pública: un boletín mensual del MEF por archivo; la hoja del indicador cambia de
# nombre entre boletines, así que se ubica por su contenido.
MES_ARCHIVO = {"ENE": 1, "FEB": 2, "MAR": 3, "ABR": 4, "MAY": 5, "JUN": 6, "JUL": 7, "AGO": 8,
               "SEP": 9, "OCT": 10, "NOV": 11, "DIC": 12}
CLAVES_DEUDA = {"externa": "total deuda externa", "interna": "total deuda interna",
                "total": "deuda pública total", "pib": "pib nominal"}


def indicador_deuda(ruta: Path) -> dict:
    import openpyxl
    wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    candidatas = [n for n in wb.sheetnames if re.fullmatch(r"1\.?", n.strip()) or "indicador" in n.lower()]
    for nombre in candidatas + [n for n in wb.sheetnames if n not in candidatas]:
        out = {}
        for r in wb[nombre].iter_rows(values_only=True):
            v = [x for x in r if x is not None]
            if len(v) < 2:
                continue
            etiqueta = limpiar(v[0]).lower()
            for k, c in CLAVES_DEUDA.items():
                if etiqueta.startswith(c) and k not in out and isinstance(v[1], (int, float)):
                    out[k] = float(v[1])
        if "total" in out and "externa" in out:
            return out
    raise ValueError(f"No se encontró el indicador de deuda en {ruta.name}")


deuda = []
for ruta in sorted((RAW / "mef_deuda").glob("*.xlsx")):
    if "PERFIL" in ruta.name.upper():
        continue
    nombre = ruta.name.upper()
    mes = next(m for clave, m in MES_ARCHIVO.items() if clave in nombre)
    anio = int(re.search(r"20\d\d", nombre)[0])
    deuda.append({"fecha": pd.Timestamp(anio, mes, 1), **indicador_deuda(ruta)})
deuda = pd.DataFrame(deuda).sort_values("fecha").drop_duplicates("fecha", keep="last").set_index("fecha")
# PIB de referencia de cada boletín (proyección del año); si falta, el del boletín más cercano.
deuda["pib"] = deuda["pib"].where(deuda["pib"] > 1e7).bfill().ffill()
FUENTE_DEUDA = "MEF, Boletín de deuda pública"
for col_, id_, nombre, corto in [
    ("externa", "deuda_externa", "Deuda pública externa", "Externa"),
    ("interna", "deuda_interna", "Deuda pública interna y otras obligaciones", "Interna"),
    ("total", "deuda_total", "Deuda pública total y otras obligaciones", "Deuda pública"),
]:
    registrar(id_, deuda.index, deuda[col_], sector="fiscal", nombre=nombre, corto=corto, unidad="USD mm",
              frecuencia="M", decimales=0, fuente=FUENTE_DEUDA, escala=1 / 1000)
registrar("deuda_pib", deuda.index, deuda["total"] / deuda["pib"] * 100, sector="fiscal",
          nombre="Deuda pública y otras obligaciones / PIB", corto="Deuda pública / PIB", unidad="% PIB",
          frecuencia="M", decimales=1, fuente=f"{FUENTE_DEUDA} (PIB de referencia de cada boletín)")

# Perfil de vencimientos (MEF): capital e intereses por año (largo plazo) y por mes (corto plazo).
perfil_lp = sorted((RAW / "mef_deuda").glob("PERFIL*LP*.xlsx"))[-1]
perfil_cp = sorted((RAW / "mef_deuda").glob("PERFIL*CP*.xlsx"))[-1]


def perfil_anual(hoja: str, prefijo: str) -> list[tuple]:
    """Totales por año de la deuda externa (por tipo de acreedor) e interna."""
    df = pd.read_excel(perfil_lp, sheet_name=hoja, header=None)
    fila_anios = next(i for i in range(20) if isinstance(df.iat[i, 1], (int, float)) and df.iat[i, 1] > 2000)
    anios = {j: int(v) for j, v in enumerate(df.iloc[fila_anios]) if isinstance(v, (int, float)) and v > 2000}
    etiquetas = [limpiar(x) if isinstance(x, str) else "" for x in df.iloc[:, 0]]
    totales = [i for i, e in enumerate(etiquetas) if e.lower() == "total general"]
    filas_ = {"externa": totales[0], "interna": totales[1]}
    if prefijo == "capital":
        for clave, texto in [("bonos", "BONOS EMITIDOS EN MERCADOS"), ("bancos", "CONVENIOS ORIGINALES (BANCOS"),
                             ("gobiernos", "CONVENIOS ORIGINALES (GOB"), ("organismos", "ORGANISMOS INTERNACIONALES")]:
            filas_[clave] = next(i for i, e in enumerate(etiquetas) if e.upper().startswith(texto))
    out = []
    for clave, i in filas_.items():
        for j, a in anios.items():
            v = pd.to_numeric(df.iat[i, j], errors="coerce")
            out.append(("anual", f"{a}-01-01", f"{prefijo}_{clave}", round((v if pd.notna(v) else 0) / 1e6, 3)))
    return out


def perfil_mensual(hoja: str, prefijo: str, rotulo: str) -> list[tuple]:
    """Total (externa + interna) de los próximos meses."""
    df = pd.read_excel(perfil_cp, sheet_name=hoja, header=None)
    i = next(i for i in range(30) if limpiar(df.iat[i, 0]).upper().startswith(f"{rotulo} DEUDA EXTERNA + INTERNA"))
    cab = df.iloc[i - 1]
    out = []
    for j in range(1, df.shape[1]):
        m = re.match(r"([A-ZÁÉÍÓÚ]+)\s+(\d{4})", limpiar(cab[j]).upper()) if isinstance(cab[j], str) else None
        if m and m[1].lower() in MESES_LARGOS:
            fecha = pd.Timestamp(int(m[2]), MESES_LARGOS[m[1].lower()], 1)
            out.append(("mensual", fecha.strftime("%Y-%m-%d"), f"{prefijo}_total",
                        round(float(pd.to_numeric(df.iat[i, j], errors="coerce") or 0) / 1e6, 3)))
    return out


vencimientos = (perfil_anual("19.", "capital") + perfil_anual("20.", "intereses")
                + perfil_mensual("16.", "capital", "CAPITAL") + perfil_mensual("17.", "intereses", "INTERES"))
pd.DataFrame(vencimientos, columns=["tipo", "fecha", "concepto", "valor"]).to_csv(OUT / "vencimientos.csv", index=False)
corte_perfil = re.search(r"(ENERO|FEBRERO|MARZO|ABRIL|MAYO|JUNIO|JULIO|AGOSTO|SEPTIEMBRE|OCTUBRE|NOVIEMBRE|DICIEMBRE)-(\d{4})",
                         perfil_lp.name.upper())
fecha_perfil = pd.Timestamp(int(corte_perfil[2]), MESES_LARGOS[corte_perfil[1].lower()], 1)

# Recaudación del SRI: hoja "Recaudación abierta" de cada archivo anual (miles de USD).
IMPUESTOS_SRI = [
    ("sri_renta", "Impuesto a la Renta", "Impuesto a la renta"),
    ("sri_iva", "Impuesto al Valor Agregado", "IVA"),
    ("sri_ice", "Impuesto a los Consumos Especiales", "ICE"),
    ("sri_isd", "Impuesto a la Salida de Divisas", "ISD"),
    ("sri_vehiculos", "Impuesto a los Vehículos Motorizados", "Vehículos"),
    ("sri_bruta", "RECAUDACIÓN BRUTA", "Recaudación bruta"),
    ("sri_neta", "RECAUDACIÓN NETA", "Recaudación neta"),
]
sri = {id_: {} for id_, _, _ in IMPUESTOS_SRI}
for ruta in sorted((RAW / "sri").glob("*.xlsx")):
    df = pd.read_excel(ruta, sheet_name="Recaudación abierta", header=None)
    fila_cab = next(i for i in range(15) if limpiar(df.iat[i, 0]).upper() == "CONCEPTOS")
    anio = int(re.search(r"20\d\d", limpiar(df.iat[1, 0]) + " " + ruta.name)[0])
    cols = {j: MESES_LARGOS[limpiar(c).lower()] for j, c in enumerate(df.iloc[fila_cab])
            if isinstance(c, str) and limpiar(c).lower() in MESES_LARGOS}
    etiquetas = [limpiar(x) if isinstance(x, str) else "" for x in df.iloc[:, 0]]
    # Meses con dato: el archivo del año en curso trae ceros en los meses futuros.
    i_neta = next(i for i, e in enumerate(etiquetas) if e.upper().startswith("RECAUDACIÓN NETA"))
    meses_con_dato = [j for j in cols if (pd.to_numeric(df.iat[i_neta, j], errors="coerce") or 0) > 0]
    for id_, texto, _ in IMPUESTOS_SRI:
        i = next(i for i, e in enumerate(etiquetas) if e.upper().startswith(texto.upper()))
        for j in meses_con_dato:
            sri[id_][pd.Timestamp(anio, cols[j], 1)] = float(pd.to_numeric(df.iat[i, j], errors="coerce") or 0) / 1000
for id_, _, corto in IMPUESTOS_SRI:
    s = pd.Series(sri[id_]).sort_index()
    registrar(id_, s.index, s, sector="fiscal", nombre=f"Recaudación SRI: {corto}", corto=corto,
              unidad="USD mm", frecuencia="M", decimales=0, fuente="SRI, Estadísticas de recaudación")

# ============================================================== SECTOR EXTERNO
bc = pd.read_excel(RAW / "balanza comercial.xlsx", header=4)
bc = bc[pd.to_datetime(bc["fecha_corte"], errors="coerce").notna()]
fechas = pd.to_datetime(bc["fecha_corte"]).dt.to_period("M").dt.to_timestamp()
for col_, id_, nombre, corto in [
    ("bc_total_exp", "x_total", "Exportaciones totales FOB", "Exportaciones"),
    ("bc_exp_pet", "x_pet", "Exportaciones petroleras FOB", "Petroleras"),
    ("bc_exp_no_pet", "x_no_pet", "Exportaciones no petroleras FOB", "No petroleras"),
    ("bc_tot_imp", "m_total", "Importaciones totales FOB", "Importaciones"),
    ("bc_imp_pet", "m_pet", "Importaciones petroleras FOB", "Petroleras"),
    ("bc_imp_no_pet", "m_no_pet", "Importaciones no petroleras FOB", "No petroleras"),
    ("bc_tot_bal_com", "bc_total", "Balanza comercial total", "Total"),
    ("bc_bal_com_pet", "bc_pet", "Balanza comercial petrolera", "Petrolera"),
    ("bc_bal_com_no_pet", "bc_no_pet", "Balanza comercial no petrolera", "No petrolera"),
]:
    registrar(id_, fechas, pd.to_numeric(bc[col_]), sector="externo", nombre=nombre, corto=corto,
              unidad="USD mm", frecuencia="M", decimales=0, fuente="BCE, Balanza Comercial")
registrar("terminos_intercambio", fechas, pd.to_numeric(bc["bc_ind_term_int"]), sector="externo",
          nombre="Índice de términos de intercambio", corto="Términos de intercambio", unidad="índice",
          frecuencia="M", decimales=1, fuente="BCE, Balanza Comercial")

bp = pd.read_excel(RAW / "Balanza de pagos.xlsx", header=None)
fila_cab = next(i for i in range(20) if limpiar(bp.iat[i, 1]) == "Componentes normalizados")
trimestres = {j: c for j, c in enumerate(bp.iloc[fila_cab]) if isinstance(c, str) and re.match(r"^(I|II|III|IV)-\d{4}$", c)}
fechas_bp = [pd.Timestamp(int(c.split("-")[1]), ROMANOS[c.split("-")[0]], 1) for c in trimestres.values()]
etiquetas = [limpiar(x) if isinstance(x, str) else "" for x in bp.iloc[:, 1]]


def fila_bp(texto: str, ocurrencia: int = 0, desde: int = 0) -> pd.Series:
    idx = [i for i, e in enumerate(etiquetas) if i >= desde and e.startswith(texto)][ocurrencia]
    return pd.to_numeric(bp.iloc[idx, list(trimestres)], errors="coerce").values


pasivos = next(i for i, e in enumerate(etiquetas) if e.startswith("Pasivos netos incurridos"))
FUENTE_BP = "BCE, Balanza de Pagos"
for id_, valores, nombre, corto in [
    ("cuenta_corriente", fila_bp("CUENTA CORRIENTE"), "Saldo de la cuenta corriente", "Cuenta corriente"),
    ("bp_bienes", fila_bp("BIENES ("), "Balanza de bienes (balanza de pagos)", "Bienes"),
    ("bp_servicios", fila_bp("SERVICIOS"), "Balanza de servicios", "Servicios"),
    ("bp_ingreso_primario", fila_bp("INGRESO PRIMARIO"), "Ingreso primario", "Ingreso primario"),
    ("bp_ingreso_secundario", fila_bp("INGRESO SECUNDARIO"), "Ingreso secundario", "Ingreso secundario"),
    ("remesas_recibidas", fila_bp("Remesas de trabajadores", 0), "Remesas recibidas", "Remesas recibidas"),
    ("remesas_enviadas", fila_bp("Remesas de trabajadores", 1), "Remesas enviadas", "Remesas enviadas"),
    ("ied", fila_bp("Inversión directa", 0, pasivos), "Inversión extranjera directa (pasivos netos)", "IED"),
]:
    registrar(id_, fechas_bp, valores, sector="externo", nombre=nombre, corto=corto, unidad="USD mm",
              frecuencia="T", decimales=0, fuente=FUENTE_BP)

cc = pd.Series(fila_bp("CUENTA CORRIENTE"), index=fechas_bp)
cc_anual = cc.groupby(cc.index.year).agg(["sum", "count"])
cc_anual = cc_anual[cc_anual["count"] == 4]["sum"]
anios = [a for a in cc_anual.index if a in pib_anual.index]
registrar("cuenta_corriente_pib", [pd.Timestamp(a, 1, 1) for a in anios],
          [cc_anual[a] / pib_anual[a] * 100 for a in anios], sector="externo",
          nombre="Saldo de la cuenta corriente", corto="Cuenta corriente", unidad="% PIB", frecuencia="A",
          decimales=1, fuente="BCE, Balanza de Pagos y Cuentas Nacionales")

# Comercio exterior por producto (exportaciones) y por uso o destino económico (importaciones)
exp = con.execute(f"""
    select make_date(year, month, 1) fecha, cast(matricero as int) codigo,
           any_value(desc_matricero) producto, sum(fob) / 1000 fob, sum(tm) / 1000 tm
    from '{RAW / "exports.parquet"}' where year >= 2018 and matricero > 0 group by 1, 2 order by 1, 2
""").df()
PETROLEO = {150101: "Crudo", 230601: "Derivados"}


def grupo_exportacion(codigo: int) -> str:
    if codigo in PETROLEO:
        return "petrolero"
    return {"1": "primario", "2": "industrializado"}.get(str(codigo)[0], "otro")


productos = (exp.groupby("codigo")["producto"].last().map(lambda p: limpiar(p).capitalize()).to_dict())
productos.update({c: n for c, n in PETROLEO.items()})
json.dump({str(c): {"producto": productos[c], "grupo": grupo_exportacion(c)} for c in sorted(productos)},
          open(OUT / "exportaciones.json", "w"), ensure_ascii=False, indent=1)
exp.assign(fecha=exp["fecha"].dt.strftime("%Y-%m-%d"), fob=exp["fob"].round(3), tm=exp["tm"].round(3))[
    ["fecha", "codigo", "fob", "tm"]].to_csv(OUT / "exportaciones.csv", index=False)

imp = con.execute(f"""
    select make_date(year, month, 1) fecha, trim(cuode1) cuode, any_value(desc_cuode1) descripcion,
           sum(fob) / 1000 fob, sum(cif) / 1000 cif
    from '{RAW / "imports.parquet"}' where year >= 2018 and trim(cuode1) not like '-%' group by 1, 2 order by 1, 2
""").df()
GRUPOS_CUODE = {"01": "Bienes de consumo", "02": "Bienes de consumo", "99": "Bienes de consumo",
                "03": "Combustibles y lubricantes", "04": "Materias primas", "05": "Materias primas",
                "06": "Materias primas", "07": "Bienes de capital", "08": "Bienes de capital",
                "09": "Bienes de capital", "10": "Diversos"}
json.dump({c: {"descripcion": limpiar(d).capitalize(), "grupo": GRUPOS_CUODE.get(c, "Diversos")}
           for c, d in imp.groupby("cuode")["descripcion"].last().items()},
          open(OUT / "importaciones.json", "w"), ensure_ascii=False, indent=1)
imp.assign(fecha=imp["fecha"].dt.strftime("%Y-%m-%d"), fob=imp["fob"].round(3), cif=imp["cif"].round(3))[
    ["fecha", "cuode", "fob", "cif"]].to_csv(OUT / "importaciones.csv", index=False)

# Control: el detalle por producto debe cuadrar con la balanza comercial.
x_detalle = exp.groupby("fecha")["fob"].sum()
x_bc = pd.Series(pd.to_numeric(bc["bc_total_exp"]).values, index=fechas)
dif = (x_detalle - x_bc.reindex(x_detalle.index)).abs().max()
print(f"Exportaciones: diferencia máxima detalle vs. balanza comercial = USD {dif:.1f} mm")

# ============================================================== MERCADOS
# Series descargadas por descargar_datos.py; el catálogo está en catalogo_mercados.py.
largo = pd.read_csv(RAW / "mercados.csv", parse_dates=["fecha"])
ancho = largo.pivot(index="fecha", columns="id", values="valor")
faltan = [id_ for id_, *_ in MERCADOS if id_ not in ancho.columns]
if faltan:
    raise ValueError(f"Sin datos de mercado para: {', '.join(faltan)} (ejecuta descargar_datos.py)")
mercados = ancho[[id_ for id_, *_ in MERCADOS]].sort_index()
mercados = mercados[mercados.index.dayofweek < 5]
meta_mercados = [dict(id=id_, sector="internacional", grupo=grupo, nombre=nombre, corto=corto, unidad=unidad,
                      frecuencia="D", decimales=dec, fuente=FUENTES[fuente], codigo=codigo)
                 for id_, fuente, codigo, nombre, corto, grupo, unidad, dec in MERCADOS]

# Las condiciones de uso de Yahoo Finance no permiten redistribuir sus datos, así que el sitio
# publica solo lo que muestra: las series que se grafican completas viajan completas; las que
# solo aparecen en tablas o curvas viajan con los valores de las fechas de referencia que usan
# las páginas (hoy, 1 mes, 3 meses, cierre del año anterior y 12 meses).
HISTORIA_COMPLETA = {"fed", "fed_inf", "bce_tasa", "boe", "boj", "us_10a", "us_2a", "dxy", "wti", "brent", "cacao",
                     "oro", "cobre", "gasolina", "diesel"}
ULTIMO_ANIO = {"spx", "stoxx50", "nikkei", "bovespa", "us_3m"}  # us_3m: el resumen lo compara con hace un año
corte = mercados.index.max()
referencias = [corte, corte - pd.Timedelta(days=30), corte - pd.Timedelta(days=91),
               pd.Timestamp(corte.year - 1, 12, 31), corte - pd.DateOffset(years=1)]
for id_ in mercados.columns:
    if id_ in HISTORIA_COMPLETA:
        continue
    s = mercados[id_].dropna()
    conservar = {s.index[s.index <= r].max() for r in referencias if (s.index <= r).any()}
    if id_ in ULTIMO_ANIO:
        conservar |= set(s.index[s.index >= referencias[-1] - pd.Timedelta(days=7)])
    mercados.loc[~mercados.index.isin(conservar), id_] = None
mercados = mercados.dropna(how="all")
mercados.index = mercados.index.strftime("%Y-%m-%d")
mercados.index.name = "fecha"
mercados.round(4).to_csv(OUT / "mercados.csv")
json.dump({"curvas": [{"id": p, "pais": n, "plazos": [t for t, _ in pl]} for p, n, _, pl in CURVAS],
           "series": meta_mercados},
          open(OUT / "mercados.json", "w"), ensure_ascii=False, indent=1)

# ============================================================== PREVISIONES DE CRECIMIENTO
# Crecimiento real del PIB previsto por institución. El BCE y el MEF llegan como Excel; el FMI
# y el Banco Mundial los descarga descargar_datos.py; otras.csv se mantiene a mano (CEPAL u
# otras instituciones), con las columnas institucion, informe, fecha, anio, valor.
PREV = RAW / "previsiones"
previsiones = [pd.read_csv(r) for r in sorted(PREV.glob("*.csv"))]


def fila_anual(ruta: Path, hoja: str, etiqueta: str) -> dict[int, float]:
    """Valores de una fila de un cuadro con años en columnas ('2026', '2026 (p)')."""
    df = pd.read_excel(ruta, sheet_name=hoja, header=None)
    i = next(i for i in range(len(df)) if any(limpiar(x).lower().startswith(etiqueta.lower())
                                              for x in df.iloc[i] if isinstance(x, str)))
    cab = next(df.iloc[k] for k in range(i - 1, -1, -1)
               if sum(bool(re.match(r"^20\d\d", str(x))) for x in df.iloc[k]) >= 2)
    return {int(str(c)[:4]): float(df.iat[i, j]) for j, c in cab.items()
            if re.match(r"^20\d\d", str(c)) and pd.notna(df.iat[i, j])}


# BCE, Previsiones macroeconómicas: PrevEconAAMM.xlsx, hoja 3 (tasas de variación anual).
ruta = mas_reciente(PREV, "PrevEcon*.xlsx")
aamm = re.search(r"(\d{2})(\d{2})", ruta.stem)
pub = pd.Timestamp(2000 + int(aamm[1]), int(aamm[2]), 1)
previsiones.append(pd.DataFrame([("BCE", f"Previsiones macroeconómicas, {list(MESES)[pub.month - 1]}-{pub.year}",
                                  pub.strftime("%Y-%m-%d"), a, v)
                                 for a, v in fila_anual(ruta, "3", "Producto Interno Bruto").items()],
                                columns=["institucion", "informe", "fecha", "anio", "valor"]))

# MEF, Programación macroeconómica: cuadro de mando de las proyecciones (solo años proyectados).
# El archivo no trae fecha de publicación; se usa la del archivo.
ruta = mas_reciente(PREV, "*Resumen-de-las-proyecciones*.xlsx")
pub = pd.Timestamp(ruta.stat().st_mtime, unit="s").to_period("M").to_timestamp()
anios = re.findall(r"20\d\d", ruta.stem)
previsiones.append(pd.DataFrame([("MEF", f"Programación macroeconómica {anios[0]}–{anios[-1]}", pub.strftime("%Y-%m-%d"), a, v)
                                 for a, v in fila_anual(ruta, "Anexo 1", "Crecimiento real del PIB").items()
                                 if a >= int(anios[0])],
                                columns=["institucion", "informe", "fecha", "anio", "valor"]))

ORDEN_INSTITUCIONES = ["BCE", "MEF", "FMI", "Banco Mundial", "CEPAL"]
previsiones = pd.concat(previsiones, ignore_index=True)
previsiones["orden"] = [ORDEN_INSTITUCIONES.index(i) if i in ORDEN_INSTITUCIONES else 99
                        for i in previsiones["institucion"]]
previsiones = (previsiones[previsiones["anio"] >= pd.Timestamp.today().year]
               .sort_values(["orden", "anio"]).drop(columns="orden"))
previsiones["valor"] = previsiones["valor"].round(2)
previsiones.to_csv(OUT / "previsiones.csv", index=False)

# ============================================================== Escritura
pd.DataFrame(filas, columns=["id", "fecha", "valor"]).to_csv(OUT / "series.csv", index=False)
json.dump({"incidencias_ipc": fecha_incidencias.strftime("%Y-%m-%d"),
           "perfil_vencimientos": fecha_perfil.strftime("%Y-%m-%d"), "series": catalogo},
          open(OUT / "catalogo.json", "w"), ensure_ascii=False, indent=1)

ids = [c["id"] for c in catalogo]
assert len(ids) == len(set(ids)), "IDs duplicados en el catálogo"
print(f"series.csv: {len(filas)} filas, {len(catalogo)} series")
print(f"mercados.csv: {mercados.shape[0]} fechas × {mercados.shape[1]} series (corte {corte:%Y-%m-%d})")
print(f"previsiones.csv: {previsiones['institucion'].nunique()} instituciones, {len(previsiones)} filas")
print(f"exportaciones.csv: {len(exp)} filas · importaciones.csv: {len(imp)} filas")
